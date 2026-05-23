from pathlib import Path

from docx import Document
from openpyxl import Workbook

from app.tools.file_reader import PdfPage, _extract_tables_from_text_pages, extract_text


def test_extract_text_from_txt(tmp_path: Path) -> None:
    path = tmp_path / "demo.txt"
    path.write_text("hello office agent", encoding="utf-8")
    assert "hello office agent" in extract_text(path)


def test_extract_text_from_docx(tmp_path: Path) -> None:
    path = tmp_path / "demo.docx"
    document = Document()
    document.add_paragraph("季度报告")
    document.save(path)
    assert "季度报告" in extract_text(path)


def test_extract_text_from_xlsx(tmp_path: Path) -> None:
    path = tmp_path / "demo.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "销售"
    sheet.append(["产品", "金额"])
    sheet.append(["A", 100])
    workbook.save(path)
    assert "销售" in extract_text(path)
    assert "产品 | 金额" in extract_text(path)


def test_extract_pdf_tables_from_text_page_fallback() -> None:
    tables = _extract_tables_from_text_pages(
        [
            PdfPage(
                page_number=3,
                text=(
                    "Table 1: Main results\n"
                    "Model  Accuracy  F1\n"
                    "Base   0.82      0.80\n"
                    "Ours   0.91      0.89\n\n"
                    "2. Related Work"
                ),
            )
        ]
    )

    assert len(tables) == 1
    assert tables[0].label == "Table 1"
    assert tables[0].page_number == 3
    assert tables[0].rows[0] == ["Model", "Accuracy", "F1"]
    assert tables[0].rows[2] == ["Ours", "0.91", "0.89"]
